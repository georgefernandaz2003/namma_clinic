import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, PageBreak, KeepTogether
)
from reportlab.pdfgen import canvas

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD_PATH = os.path.join(BASE_DIR, "docs", "USER_DOCUMENT.md")
DOCX_OUT = os.path.join(BASE_DIR, "Namma_Clinic_User_Manual.docx")
DOCX_OUT_DOCS = os.path.join(BASE_DIR, "docs", "Namma_Clinic_User_Manual.docx")
PDF_OUT = os.path.join(BASE_DIR, "Namma_Clinic_User_Manual.pdf")
PDF_OUT_DOCS = os.path.join(BASE_DIR, "docs", "Namma_Clinic_User_Manual.pdf")

def read_markdown_file(path):
    with open(path, 'r', encoding='utf-8') as f:
        return f.read()

# ==================== DOCX BUILDER ====================
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_docx(md_content):
    doc = Document()
    
    # Configure margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Document Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("🏥 Namma Clinic Digital Healthcare Network")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(6, 95, 70) # Deep Emerald

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("Comprehensive User Operations Manual & Clinical Workflow Guide\nVersion 2.4.0 • September 2026")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(100, 116, 139) # Slate

    doc.add_paragraph() # Spacer

    lines = md_content.split('\n')
    i = 0
    in_table = False
    table_rows = []
    in_alert = False
    alert_lines = []

    while i < len(lines):
        line = lines[i].strip()
        
        # Skip top markdown header since we added styled title
        if line.startswith("# 📖 Namma Clinic Digital Healthcare Network"):
            i += 1
            continue

        # Handle Alert box (> [!TIP], > [!NOTE], etc.)
        if line.startswith('> [!') or (in_alert and line.startswith('>')):
            in_alert = True
            clean_l = line.replace('> [!TIP]', '💡 TIP:').replace('> [!NOTE]', '📌 NOTE:').replace('> [!IMPORTANT]', '⚠️ IMPORTANT:').lstrip('> ').strip()
            if clean_l:
                alert_lines.append(clean_l)
            i += 1
            continue
        elif in_alert:
            # End of alert
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            set_cell_background(cell, "F0FDF4") # Light emerald tint
            set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
            p = cell.paragraphs[0]
            for al in alert_lines:
                r = p.add_run(al + "\n")
                r.font.name = "Arial"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(22, 101, 52)
            doc.add_paragraph()
            in_alert = False
            alert_lines = []

        # Handle Markdown Table (| ... |)
        if line.startswith('|') and line.endswith('|'):
            # Check if divider row
            if re.match(r'^\|[\s\-:]+\|$', line):
                i += 1
                continue
            cells = [c.strip() for c in line.split('|')[1:-1]]
            table_rows.append(cells)
            i += 1
            in_table = True
            continue
        elif in_table:
            # Render accumulated table
            if table_rows:
                num_cols = max(len(r) for r in table_rows)
                tbl = doc.add_table(rows=len(table_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row_data in enumerate(table_rows):
                    for c_idx in range(num_cols):
                        val = row_data[c_idx] if c_idx < len(row_data) else ""
                        c = tbl.cell(r_idx, c_idx)
                        set_cell_margins(c, top=80, bottom=80, left=120, right=120)
                        cp = c.paragraphs[0]
                        # Remove markdown bold tags
                        clean_val = val.replace('**', '')
                        c_run = cp.add_run(clean_val)
                        c_run.font.name = "Arial"
                        c_run.font.size = Pt(9)
                        if r_idx == 0:
                            set_cell_background(c, "E2E8F0") # Slate header
                            c_run.font.bold = True
                            c_run.font.color.rgb = RGBColor(15, 23, 42)
                        else:
                            if r_idx % 2 == 1:
                                set_cell_background(c, "F8FAFC")
                            c_run.font.color.rgb = RGBColor(51, 65, 85)
                doc.add_paragraph()
            table_rows = []
            in_table = False

        # Headings
        if line.startswith('## '):
            h = doc.add_heading(level=1)
            r = h.add_run(line.replace('## ', '').strip())
            r.font.name = "Arial"
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
        elif line.startswith('### '):
            h = doc.add_heading(level=2)
            r = h.add_run(line.replace('### ', '').strip())
            r.font.name = "Arial"
            r.font.size = Pt(11.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(30, 41, 59)
        elif line.startswith('#### '):
            h = doc.add_heading(level=3)
            r = h.add_run(line.replace('#### ', '').strip())
            r.font.name = "Arial"
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(51, 65, 85)
        # Bullet list
        elif line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            clean_text = line[2:].strip()
            # Process simple bold tags
            parts = re.split(r'(\*\*.*?\*\*)', clean_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                else:
                    run = p.add_run(part)
                run.font.name = "Arial"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(51, 65, 85)
        # Numbered list
        elif re.match(r'^\d+\.\s', line):
            p = doc.add_paragraph(style='List Number')
            clean_text = re.sub(r'^\d+\.\s', '', line).strip()
            parts = re.split(r'(\*\*.*?\*\*)', clean_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                else:
                    run = p.add_run(part)
                run.font.name = "Arial"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(51, 65, 85)
        # Code block (```)
        elif line.startswith('```'):
            code_block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_block.append(lines[i])
                i += 1
            code_str = "\n".join(code_block)
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            c = tbl.cell(0, 0)
            set_cell_background(c, "F1F5F9")
            set_cell_margins(c, top=100, bottom=100, left=150, right=150)
            cp = c.paragraphs[0]
            cr = cp.add_run(code_str)
            cr.font.name = "Consolas"
            cr.font.size = Pt(8.5)
            cr.font.color.rgb = RGBColor(15, 23, 42)
            doc.add_paragraph()
        elif line == '---':
            p = doc.add_paragraph()
            r = p.add_run("_______________________________________________________________________________")
            r.font.color.rgb = RGBColor(203, 213, 225)
            r.font.size = Pt(8)
        elif line:
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                else:
                    run = p.add_run(part)
                run.font.name = "Arial"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(51, 65, 85)
        
        i += 1

    doc.save(DOCX_OUT)
    doc.save(DOCX_OUT_DOCS)
    print(f"Word document saved to: {DOCX_OUT}")


# ==================== REPORTLAB PDF BUILDER ====================
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_header_footer(num_pages)
            super().showPage()
        super().save()

    def draw_header_footer(self, page_count):
        self.saveState()
        if self._pageNumber > 1:
            # Header
            self.setFont("Helvetica-Bold", 8)
            self.setFillColor(colors.HexColor("#065F46"))
            self.drawString(36, 762, "NAMMA CLINIC DIGITAL HEALTHCARE NETWORK")
            self.setFont("Helvetica", 8)
            self.setFillColor(colors.HexColor("#64748B"))
            self.drawRightString(576, 762, "User Manual & Operations Guide • v2.4.0")
            self.setStrokeColor(colors.HexColor("#CBD5E1"))
            self.setLineWidth(0.5)
            self.line(36, 756, 576, 756)

            # Footer
            self.line(36, 36, 576, 36)
            self.setFont("Helvetica", 8)
            self.setFillColor(colors.HexColor("#64748B"))
            self.drawString(36, 24, "Urban Primary Healthcare Network • Department of Health & Family Welfare")
            page_text = f"Page {self._pageNumber} of {page_count}"
            self.drawRightString(576, 24, page_text)
        self.restoreState()

def clean_for_pdf(text):
    text = text.replace('&', '&amp;')
    text = text.replace('<', '&lt;').replace('>', '&gt;')
    # restore basic formatting
    text = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', text)
    text = re.sub(r'\*(.*?)\*', r'<i>\1</i>', text)
    text = re.sub(r'`(.*?)`', r'<font face="Courier">\1</font>', text)
    return text

def generate_pdf(md_content):
    doc = SimpleDocTemplate(
        PDF_OUT,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#065F46'),
        alignment=1, # Center
        spaceAfter=6
    )

    sub_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10.5,
        leading=14,
        textColor=colors.HexColor('#64748B'),
        alignment=1,
        spaceAfter=15
    )

    h1_style = ParagraphStyle(
        'Heading1_Custom',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=17,
        textColor=colors.HexColor('#0F172A'),
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True
    )

    h2_style = ParagraphStyle(
        'Heading2_Custom',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#1E293B'),
        spaceBefore=8,
        spaceAfter=4,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        'Body_Custom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12.5,
        textColor=colors.HexColor('#334155'),
        spaceAfter=4
    )

    bullet_style = ParagraphStyle(
        'Bullet_Custom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12.5,
        textColor=colors.HexColor('#334155'),
        leftIndent=14,
        firstLineIndent=-10,
        spaceAfter=3
    )

    code_style = ParagraphStyle(
        'Code_Custom',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor('#0F172A')
    )

    table_cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor('#1E293B')
    )

    table_hdr_style = ParagraphStyle(
        'TableHdr',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor('#FFFFFF')
    )

    alert_style = ParagraphStyle(
        'AlertText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor('#166534')
    )

    elements = []

    # Title block
    elements.append(Paragraph("🏥 Namma Clinic Digital Healthcare Network", title_style))
    elements.append(Paragraph("User Operations Manual & Clinical Workflow Guide<br/><b>Version 2.4.0 • Published September 2026</b>", sub_style))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#065F46'), spaceAfter=12))

    lines = md_content.split('\n')
    i = 0
    in_table = False
    table_rows = []
    in_alert = False
    alert_lines = []

    while i < len(lines):
        line = lines[i].strip()

        if line.startswith("# 📖 Namma Clinic"):
            i += 1
            continue

        # Handle Alert box
        if line.startswith('> [!') or (in_alert and line.startswith('>')):
            in_alert = True
            clean_l = line.replace('> [!TIP]', '💡 <b>TIP:</b>').replace('> [!NOTE]', '📌 <b>NOTE:</b>').replace('> [!IMPORTANT]', '⚠️ <b>IMPORTANT:</b>').lstrip('> ').strip()
            if clean_l:
                alert_lines.append(clean_for_pdf(clean_l))
            i += 1
            continue
        elif in_alert:
            alert_text = "<br/>".join(alert_lines)
            t = Table([[Paragraph(alert_text, alert_style)]], colWidths=[540])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F0FDF4')),
                ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#86EFAC')),
                ('LEFTPADDING', (0,0), (-1,-1), 10),
                ('RIGHTPADDING', (0,0), (-1,-1), 10),
                ('TOPPADDING', (0,0), (-1,-1), 8),
                ('BOTTOMPADDING', (0,0), (-1,-1), 8),
            ]))
            elements.append(Spacer(1, 4))
            elements.append(t)
            elements.append(Spacer(1, 6))
            in_alert = False
            alert_lines = []

        # Handle Markdown Table
        if line.startswith('|') and line.endswith('|'):
            if re.match(r'^\|[\s\-:]+\|$', line):
                i += 1
                continue
            cells = [c.strip() for c in line.split('|')[1:-1]]
            table_rows.append(cells)
            i += 1
            in_table = True
            continue
        elif in_table:
            if table_rows:
                num_cols = max(len(r) for r in table_rows)
                col_width = 540 / num_cols
                pdf_table_data = []
                for r_idx, r_data in enumerate(table_rows):
                    row_cells = []
                    for c_idx in range(num_cols):
                        val = r_data[c_idx] if c_idx < len(r_data) else ""
                        c_text = clean_for_pdf(val)
                        if r_idx == 0:
                            row_cells.append(Paragraph(c_text, table_hdr_style))
                        else:
                            row_cells.append(Paragraph(c_text, table_cell_style))
                    pdf_table_data.append(row_cells)

                t = Table(pdf_table_data, colWidths=[col_width] * num_cols)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#065F46')),
                    ('TEXTCOLOR', (0,0), (-1,0), colors.white),
                    ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
                    ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F8FAFC')]),
                    ('TOPPADDING', (0,0), (-1,-1), 5),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 5),
                    ('LEFTPADDING', (0,0), (-1,-1), 6),
                    ('RIGHTPADDING', (0,0), (-1,-1), 6),
                ]))
                elements.append(Spacer(1, 4))
                elements.append(t)
                elements.append(Spacer(1, 6))
            table_rows = []
            in_table = False

        # Headings
        if line.startswith('## '):
            text = clean_for_pdf(line.replace('## ', '').strip())
            elements.append(Spacer(1, 6))
            elements.append(Paragraph(text, h1_style))
            elements.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#E2E8F0'), spaceAfter=4))
        elif line.startswith('### '):
            text = clean_for_pdf(line.replace('### ', '').strip())
            elements.append(Paragraph(text, h2_style))
        elif line.startswith('#### '):
            text = clean_for_pdf(line.replace('#### ', '').strip())
            elements.append(Paragraph(f"<b>{text}</b>", body_style))
        elif line.startswith('- ') or line.startswith('* '):
            text = clean_for_pdf(line[2:].strip())
            elements.append(Paragraph(f"&bull; {text}", bullet_style))
        elif re.match(r'^\d+\.\s', line):
            text = clean_for_pdf(line.strip())
            elements.append(Paragraph(text, bullet_style))
        elif line.startswith('```'):
            code_block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_block.append(lines[i].replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;'))
                i += 1
            code_str = "<br/>".join(code_block)
            t = Table([[Paragraph(code_str, code_style)]], colWidths=[540])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F1F5F9')),
                ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
                ('LEFTPADDING', (0,0), (-1,-1), 8),
                ('RIGHTPADDING', (0,0), (-1,-1), 8),
                ('TOPPADDING', (0,0), (-1,-1), 6),
                ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ]))
            elements.append(Spacer(1, 3))
            elements.append(t)
            elements.append(Spacer(1, 5))
        elif line == '---':
            elements.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#E2E8F0'), spaceAfter=6, spaceBefore=4))
        elif line:
            text = clean_for_pdf(line)
            elements.append(Paragraph(text, body_style))

        i += 1

    doc.build(elements, canvasmaker=NumberedCanvas)
    
    # Also copy to docs/
    with open(PDF_OUT, 'rb') as f_src, open(PDF_OUT_DOCS, 'wb') as f_dst:
        f_dst.write(f_src.read())
        
    print(f"PDF document saved to: {PDF_OUT}")

if __name__ == '__main__':
    md = read_markdown_file(MD_PATH)
    print("Generating Word (.docx) document...")
    generate_docx(md)
    print("Generating PDF (.pdf) document...")
    generate_pdf(md)
    print("Successfully built both User Documents!")
